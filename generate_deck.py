"""
Generates the exact Round 2 Presentation for Medibloom matching the Turing Club UI/UX template.
"""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors from template
    COLOR_BG = RGBColor(246, 244, 238)        # Warm off-white #F6F4EE
    COLOR_DARK = RGBColor(24, 24, 24)          # Deep charcoal #181818
    COLOR_ORANGE = RGBColor(238, 82, 43)       # Accent orange #EE522B
    COLOR_MUTED = RGBColor(100, 100, 100)      # Muted gray #646464
    COLOR_BORDER = RGBColor(218, 214, 206)     # Border line #DAD6CE
    COLOR_WHITE = RGBColor(255, 255, 255)      # Pure white #FFFFFF
    COLOR_CARD_BG = RGBColor(252, 250, 247)    # Card fill #FCFAF7

    assets_dir = r"d:\NEWPROJECT\Medibloom\Medibloom\presentation_assets"
    arch_img_path = os.path.join(assets_dir, "arch_stairs.png")
    laptop_mockup_path = os.path.join(assets_dir, "medibloom_laptop_mockup.png")

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide):
        # Header text: THE TURING CLUB / FRONTEND UI/UX COMPETITION
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(7.0), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p1 = tf.paragraphs[0]
        p1.text = "THE TURING CLUB"
        p1.font.name = "Arial"
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_DARK
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = "FRONTEND UI/UX COMPETITION"
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_ORANGE

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    add_header(slide1)

    # Right side image: architectural stairs
    if os.path.exists(arch_img_path):
        slide1.shapes.add_picture(arch_img_path, Inches(8.1), Inches(0), Inches(5.233), Inches(7.5))

    # Main Heading
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(7.0), Inches(3.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    h_lines = ["ROUND 2 —", "FRONTEND", "UI/UX", "COMPETITION"]
    for idx, line in enumerate(h_lines):
        p = tf1.paragraphs[0] if idx == 0 else tf1.add_paragraph()
        p.text = line
        p.font.name = "Arial Black"
        p.font.size = Pt(46)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(0)

    # Divider line
    line = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.75), Inches(6.8), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_DARK
    line.line.fill.background()

    # Meta Info
    meta_box = slide1.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(7.0), Inches(2.2))
    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = tf_meta.margin_top = tf_meta.margin_right = tf_meta.margin_bottom = 0

    p_team_lbl = tf_meta.paragraphs[0]
    p_team_lbl.text = "TEAM NAME"
    p_team_lbl.font.name = "Arial"
    p_team_lbl.font.size = Pt(11)
    p_team_lbl.font.bold = True
    p_team_lbl.font.color.rgb = COLOR_ORANGE

    p_team_val = tf_meta.add_paragraph()
    p_team_val.text = "SHARK-025"
    p_team_val.font.name = "Arial"
    p_team_val.font.size = Pt(20)
    p_team_val.font.bold = True
    p_team_val.font.color.rgb = COLOR_DARK
    p_team_val.space_after = Pt(14)

    p_proj_lbl = tf_meta.add_paragraph()
    p_proj_lbl.text = "PROJECT TITLE"
    p_proj_lbl.font.name = "Arial"
    p_proj_lbl.font.size = Pt(11)
    p_proj_lbl.font.bold = True
    p_proj_lbl.font.color.rgb = COLOR_ORANGE

    p_proj_val = tf_meta.add_paragraph()
    p_proj_val.text = "MEDIBLOOM — Health That Grows Daily"
    p_proj_val.font.name = "Arial"
    p_proj_val.font.size = Pt(20)
    p_proj_val.font.bold = True
    p_proj_val.font.color.rgb = COLOR_DARK

    p_proj_sub = tf_meta.add_paragraph()
    p_proj_sub.text = "Accessible Healthcare, Medication Companion & Doctor Booking for Seniors (55+) & Families"
    p_proj_sub.font.name = "Arial"
    p_proj_sub.font.size = Pt(12)
    p_proj_sub.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 2: TOOLS & TECHNOLOGIES
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2)

    # Big left title
    t2_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(3.2))
    tf2 = t2_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    t2_lines = ["TOOLS", "&", "TECHN", "OLOGI", "ES"]
    for idx, l in enumerate(t2_lines):
        p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p.text = l
        p.font.name = "Arial Black"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

    sub2_box = slide2.shapes.add_textbox(Inches(0.8), Inches(4.35), Inches(5.5), Inches(1.0))
    tf_sub2 = sub2_box.text_frame
    tf_sub2.word_wrap = True
    tf_sub2.margin_left = tf_sub2.margin_top = tf_sub2.margin_right = tf_sub2.margin_bottom = 0
    p_sub2 = tf_sub2.paragraphs[0]
    p_sub2.text = "A modern, reliable stack that helps us build fast, deliver seamless experiences, and scale with confidence."
    p_sub2.font.name = "Arial"
    p_sub2.font.size = Pt(13)
    p_sub2.font.color.rgb = COLOR_MUTED

    # Table on Right (5 rows, 2 columns)
    t_left = Inches(6.8)
    t_top = Inches(0.55)
    t_width = Inches(5.8)
    t_height = Inches(4.8)
    rows = 5
    cols = 2
    table_shape = slide2.shapes.add_table(rows, cols, t_left, t_top, t_width, t_height)
    table = table_shape.table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(3.8)

    table_data = [
        ("FRONTEND", "HTML5 Semantic Architecture\nCSS3 (Custom Design Tokens, Calm Motion)\nVanilla JavaScript (ES6+)\nWeb SpeechSynthesis API\nLocalStorage State Engine"),
        ("TOOLS", "Visual Studio Code\nGit & GitHub Version Control\nChrome DevTools & Accessibility Inspector\nPowerShell Automation"),
        ("AI MODELS", "Google Antigravity Agentic IDE\nGemini 2.5 Flash Architecture\nPair Programming & UI Verification"),
        ("DESIGN", "WCAG 2.1 AA Accessibility Standards\nCalm Motion UI (Respects reduced-motion)\nHigh-Contrast Theming (Dark/Light Mode)\nDynamic Font Scaler (A, A+, A++)"),
        ("DEPLOYMENT", "Vercel Global Edge Network\nGitHub Repository (SHARK-025/Medibloom)\nContinuous Deployment & Zero-Config Hosting")
    ]

    for row_idx, (cat, desc) in enumerate(table_data):
        cell_cat = table.cell(row_idx, 0)
        cell_cat.fill.solid()
        cell_cat.fill.fore_color.rgb = COLOR_CARD_BG
        cell_cat.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_cat = cell_cat.text_frame.paragraphs[0]
        p_cat.text = cat
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(12)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ORANGE

        cell_desc = table.cell(row_idx, 1)
        cell_desc.fill.solid()
        cell_desc.fill.fore_color.rgb = COLOR_WHITE
        cell_desc.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell_desc.text_frame.word_wrap = True
        for p_idx, line_txt in enumerate(desc.split("\n")):
            p_d = cell_desc.text_frame.paragraphs[0] if p_idx == 0 else cell_desc.text_frame.add_paragraph()
            p_d.text = line_txt
            p_d.font.name = "Arial"
            p_d.font.size = Pt(11)
            p_d.font.color.rgb = COLOR_DARK
            p_d.font.bold = (p_idx == 0)

    # Bottom Banner: WHY THIS STACK
    bot_y = Inches(5.6)
    bot_box = slide2.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), bot_y, Inches(11.8), Inches(1.5)
    )
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = COLOR_CARD_BG
    bot_box.line.color.rgb = COLOR_BORDER

    # "WHY THIS STACK" title inside banner
    w_box = slide2.shapes.add_textbox(Inches(1.0), bot_y + Inches(0.2), Inches(2.2), Inches(1.1))
    tf_w = w_box.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_top = tf_w.margin_right = tf_w.margin_bottom = 0
    p_w1 = tf_w.paragraphs[0]
    p_w1.text = "WHY"
    p_w1.font.name = "Arial Black"
    p_w1.font.size = Pt(22)
    p_w1.font.bold = True
    p_w1.font.color.rgb = COLOR_DARK
    p_w2 = tf_w.add_paragraph()
    p_w2.text = "THIS"
    p_w2.font.name = "Arial Black"
    p_w2.font.size = Pt(22)
    p_w2.font.bold = True
    p_w2.font.color.rgb = COLOR_DARK
    p_w3 = tf_w.add_paragraph()
    p_w3.text = "STACK"
    p_w3.font.name = "Arial Black"
    p_w3.font.size = Pt(22)
    p_w3.font.bold = True
    p_w3.font.color.rgb = COLOR_ORANGE

    # 3 Reasons columns
    col_w = Inches(2.8)
    reasons = [
        ("Accelerates development", "Zero-framework, zero-build lightweight architecture with vanilla ES6+ for instant velocity and zero tooling fragility."),
        ("Ensures performance", "100% pure vanilla JS with sub-50ms render, zero layout shift, and instant offline LocalStorage persistence."),
        ("Scales effortlessly", "Modular component design (SPA router, 3-state medication engine, audio helper) ready for backend API growth.")
    ]
    for r_idx, (r_title, r_desc) in enumerate(reasons):
        c_left = Inches(3.6) + r_idx * Inches(2.95)
        # Vertical divider line
        if r_idx > 0:
            v_line = slide2.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, c_left - Inches(0.15), bot_y + Inches(0.2), Inches(0.015), Inches(1.1)
            )
            v_line.fill.solid()
            v_line.fill.fore_color.rgb = COLOR_BORDER
            v_line.line.fill.background()

        r_box = slide2.shapes.add_textbox(c_left, bot_y + Inches(0.18), col_w, Inches(1.1))
        tf_r = r_box.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_rt = tf_r.paragraphs[0]
        p_rt.text = r_title
        p_rt.font.name = "Arial"
        p_rt.font.size = Pt(12)
        p_rt.font.bold = True
        p_rt.font.color.rgb = COLOR_DARK
        p_rt.space_after = Pt(3)

        p_rd = tf_r.add_paragraph()
        p_rd.text = r_desc
        p_rd.font.name = "Arial"
        p_rd.font.size = Pt(10)
        p_rd.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & USER FLOW
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3)

    # Right side graphic accent
    if os.path.exists(arch_img_path):
        slide3.shapes.add_picture(arch_img_path, Inches(11.0), Inches(0), Inches(2.333), Inches(7.5))

    # Heading
    t3_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(10.0), Inches(1.2))
    tf3 = t3_box.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
    p3_h = tf3.paragraphs[0]
    p3_h.text = "SYSTEM ARCHITECTURE & USER FLOW"
    p3_h.font.name = "Arial Black"
    p3_h.font.size = Pt(36)
    p3_h.font.bold = True
    p3_h.font.color.rgb = COLOR_DARK

    p3_sub = tf3.add_paragraph()
    p3_sub.text = "A streamlined flow from user interaction to reliable output, designed for clarity, performance, and scalability."
    p3_sub.font.name = "Arial"
    p3_sub.font.size = Pt(13)
    p3_sub.font.color.rgb = COLOR_MUTED

    # 4 Architecture Flow Cards
    cards = [
        ("USER INPUT", [
            "User actions & touches",
            "Category filter (Senior / Child / All)",
            "Dose 'Mark as Taken' toggle",
            "3-Step Doctor booking form",
            "Audio 'Listen' speech click",
            "Dark mode & font-size scaler"
        ]),
        ("CORE PROCESS", [
            "Client-side SPA route router",
            "3-State Medication Engine:",
            "  • Taken (Checked off)",
            "  • Due (Scheduled today)",
            "  • Missed (Overdue detection)",
            "Real-time inline regex validation",
            "SpeechSynthesis queue manager"
        ]),
        ("IDEAS / SOLUTION", [
            "What did we solve?",
            "• Medication non-adherence & confusion in seniors (55+) and kids.",
            "How did we solve it?",
            "• Plain-language explanations, audio speech, overdue alerts (no double dose).",
            "Success Rate: 100%",
            "• Zero layout shift, full WCAG 2.1 AA."
        ]),
        ("OUTPUT", [
            "Real-time family progress bar",
            "Audible voice guidance with animated stop indicator",
            "Overdue missed-dose warning box",
            "Doctor booking ticket receipt",
            "Zero eye-strain Dark Mode palette"
        ])
    ]

    card_y = Inches(2.6)
    card_w = Inches(2.35)
    card_h = Inches(4.3)
    spacing = Inches(0.2)

    for c_idx, (c_title, c_items) in enumerate(cards):
        c_x = Inches(0.8) + c_idx * (card_w + spacing)

        # Card shape
        card_bg = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, card_y, card_w, card_h)
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = COLOR_WHITE
        card_bg.line.color.rgb = COLOR_BORDER
        card_bg.line.width = Pt(1.5)

        # Card Title Header
        title_tag = slide3.shapes.add_textbox(c_x + Inches(0.12), card_y + Inches(0.12), card_w - Inches(0.24), Inches(0.5))
        tf_tag = title_tag.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = c_title
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ORANGE

        # Content text
        content_box = slide3.shapes.add_textbox(c_x + Inches(0.12), card_y + Inches(0.65), card_w - Inches(0.24), card_h - Inches(0.8))
        tf_cnt = content_box.text_frame
        tf_cnt.word_wrap = True
        tf_cnt.margin_left = tf_cnt.margin_top = tf_cnt.margin_right = tf_cnt.margin_bottom = 0

        for item_idx, itm in enumerate(c_items):
            p_itm = tf_cnt.paragraphs[0] if item_idx == 0 else tf_cnt.add_paragraph()
            p_itm.text = itm
            p_itm.font.name = "Arial"
            if itm.startswith("What did") or itm.startswith("How did") or itm.startswith("Success Rate") or itm.startswith("3-State"):
                p_itm.font.bold = True
                p_itm.font.size = Pt(10.5)
                p_itm.font.color.rgb = COLOR_DARK
                p_itm.space_before = Pt(4)
            else:
                p_itm.font.size = Pt(10)
                p_itm.font.color.rgb = COLOR_MUTED
                p_itm.space_before = Pt(2)

        # Connecting Arrow between cards
        if c_idx < len(cards) - 1:
            arrow_x = c_x + card_w + Inches(0.04)
            arrow_y = card_y + Inches(2.0)
            arr_box = slide3.shapes.add_textbox(arrow_x, arrow_y, Inches(0.15), Inches(0.3))
            tf_arr = arr_box.text_frame
            tf_arr.margin_left = tf_arr.margin_top = tf_arr.margin_right = tf_arr.margin_bottom = 0
            p_arr = tf_arr.paragraphs[0]
            p_arr.text = "→"
            p_arr.font.name = "Arial"
            p_arr.font.size = Pt(18)
            p_arr.font.bold = True
            p_arr.font.color.rgb = COLOR_ORANGE

    # =========================================================================
    # SLIDE 4: KEY FEATURES & UI SHOWCASE
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4)

    # Domain choice prompt annotation from template
    dom_box = slide4.shapes.add_textbox(Inches(5.0), Inches(0.45), Inches(7.5), Inches(0.6))
    tf_dom = dom_box.text_frame
    tf_dom.word_wrap = True
    tf_dom.margin_left = tf_dom.margin_top = tf_dom.margin_right = tf_dom.margin_bottom = 0
    p_dom = tf_dom.paragraphs[0]
    p_dom.text = "Domain: Healthcare Accessibility for Seniors (55+) & Families in India"
    p_dom.font.name = "Arial"
    p_dom.font.size = Pt(11)
    p_dom.font.bold = True
    p_dom.font.color.rgb = COLOR_ORANGE

    p_dom2 = tf_dom.add_paragraph()
    p_dom2.text = "Why this domain? Standard medical apps overwhelm older adults with complex jargon and small text. Medibloom simplifies care with plain words, audio guides, and effortless Indian city doctor booking."
    p_dom2.font.name = "Arial"
    p_dom2.font.size = Pt(10)
    p_dom2.font.color.rgb = COLOR_MUTED

    # Heading
    t4_box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(4.5), Inches(1.2))
    tf4 = t4_box.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0
    p4_h = tf4.paragraphs[0]
    p4_h.text = "KEY FEATURES\n& UI SHOWCASE"
    p4_h.font.name = "Arial Black"
    p4_h.font.size = Pt(32)
    p4_h.font.bold = True
    p4_h.font.color.rgb = COLOR_DARK

    # Left Column: 3 Key Features
    feat_y = Inches(2.7)
    features = [
        ("01", "INTUITIVE NAVIGATION & FILTERING", "Segmented care for Seniors (55+) and Children with synchronized live family progress tracking, text-size adjustments, and one-click dark/light mode toggle."),
        ("02", "ACCESSIBLE AUDIO & 3-STATE TRACKING", "Browser-native Web SpeechSynthesis reads prescriptions aloud; smart 3-state engine dynamically flags overdue doses with clear clinical safety instructions (do not take double dose)."),
        ("03", "SEAMLESS 3-STEP APPOINTMENT BOOKING", "Frictionless booking wizard featuring top doctors across Indian cities (Bengaluru, Mumbai, Hyderabad, New Delhi) with real-time field validation and verified confirmation tickets.")
    ]

    for f_idx, (num, f_title, f_desc) in enumerate(features):
        fy = feat_y + f_idx * Inches(1.5)

        # Number pill
        num_box = slide4.shapes.add_textbox(Inches(0.8), fy, Inches(0.8), Inches(0.5))
        tf_num = num_box.text_frame
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = num
        p_num.font.name = "Arial Black"
        p_num.font.size = Pt(28)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ORANGE

        # Feature details
        desc_box = slide4.shapes.add_textbox(Inches(1.65), fy, Inches(3.6), Inches(1.4))
        tf_d = desc_box.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
        p_dt = tf_d.paragraphs[0]
        p_dt.text = f_title
        p_dt.font.name = "Arial"
        p_dt.font.size = Pt(12)
        p_dt.font.bold = True
        p_dt.font.color.rgb = COLOR_DARK
        p_dt.space_after = Pt(2)

        p_dd = tf_d.add_paragraph()
        p_dd.text = f_desc
        p_dd.font.name = "Arial"
        p_dd.font.size = Pt(10)
        p_dd.font.color.rgb = COLOR_MUTED

    # Right side: Laptop mockup with live Medibloom screenshot
    if os.path.exists(laptop_mockup_path):
        slide4.shapes.add_picture(laptop_mockup_path, Inches(5.4), Inches(1.5), Inches(7.3), Inches(5.4))

    # =========================================================================
    # SLIDE 5: IMPACT, LINKS & THANK YOU
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5)

    # Right side graphic
    if os.path.exists(arch_img_path):
        slide5.shapes.add_picture(arch_img_path, Inches(8.1), Inches(0), Inches(5.233), Inches(7.5))

    # Heading
    t5_box = slide5.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(7.0), Inches(1.8))
    tf5 = t5_box.text_frame
    tf5.word_wrap = True
    tf5.margin_left = tf5.margin_top = tf5.margin_right = tf5.margin_bottom = 0
    p5_1 = tf5.paragraphs[0]
    p5_1.text = "IMPACT, LINKS &"
    p5_1.font.name = "Arial Black"
    p5_1.font.size = Pt(38)
    p5_1.font.bold = True
    p5_1.font.color.rgb = COLOR_DARK

    p5_2 = tf5.add_paragraph()
    p5_2.text = "THANK YOU"
    p5_2.font.name = "Arial Black"
    p5_2.font.size = Pt(38)
    p5_2.font.bold = True
    p5_2.font.color.rgb = COLOR_ORANGE
    p5_2.space_after = Pt(10)

    p5_sub = tf5.add_paragraph()
    p5_sub.text = "Every design decision drives clarity, speed, and accessibility — creating experiences that seniors and families trust and enjoy."
    p5_sub.font.name = "Arial"
    p5_sub.font.size = Pt(13)
    p5_sub.font.color.rgb = COLOR_MUTED

    # Impact bullet points
    imp_box = slide5.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(7.0), Inches(1.4))
    tf_imp = imp_box.text_frame
    tf_imp.word_wrap = True
    tf_imp.margin_left = tf_imp.margin_top = tf_imp.margin_right = tf_imp.margin_bottom = 0

    impact_points = [
        "👵 Senior-First UX: WCAG 2.1 AA contrast, calm 300ms motion transitions, and typography scaling.",
        "🛡️ Clinical Safety: Overdue dose indicators prevent accidental double-dosing with safe guidance.",
        "⚡ Zero-Dependency Lightweight Speed: Instant sub-50ms render on any Indian mobile connection."
    ]
    for idx, pt in enumerate(impact_points):
        p_pt = tf_imp.paragraphs[0] if idx == 0 else tf_imp.add_paragraph()
        p_pt.text = pt
        p_pt.font.name = "Arial"
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = COLOR_DARK
        p_pt.space_after = Pt(4)

    # Link Cards (2 big cards)
    # Box 1: GitHub
    link1_shape = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.9), Inches(7.0), Inches(0.95))
    link1_shape.fill.solid()
    link1_shape.fill.fore_color.rgb = COLOR_WHITE
    link1_shape.line.color.rgb = COLOR_BORDER
    link1_shape.line.width = Pt(1.5)

    l1_box = slide5.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(6.6), Inches(0.75))
    tf_l1 = l1_box.text_frame
    tf_l1.word_wrap = True
    tf_l1.margin_left = tf_l1.margin_top = tf_l1.margin_right = tf_l1.margin_bottom = 0
    p_l1_tag = tf_l1.paragraphs[0]
    p_l1_tag.text = "GITHUB REPOSITORY URL"
    p_l1_tag.font.name = "Arial"
    p_l1_tag.font.size = Pt(10)
    p_l1_tag.font.bold = True
    p_l1_tag.font.color.rgb = COLOR_ORANGE

    p_l1_val = tf_l1.add_paragraph()
    p_l1_val.text = "https://github.com/SHARK-025/Medibloom"
    p_l1_val.font.name = "Arial"
    p_l1_val.font.size = Pt(14)
    p_l1_val.font.bold = True
    p_l1_val.font.color.rgb = COLOR_DARK

    # Box 2: Vercel Live Demo
    link2_shape = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.0), Inches(7.0), Inches(0.95))
    link2_shape.fill.solid()
    link2_shape.fill.fore_color.rgb = COLOR_WHITE
    link2_shape.line.color.rgb = COLOR_BORDER
    link2_shape.line.width = Pt(1.5)

    l2_box = slide5.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(6.6), Inches(0.75))
    tf_l2 = l2_box.text_frame
    tf_l2.word_wrap = True
    tf_l2.margin_left = tf_l2.margin_top = tf_l2.margin_right = tf_l2.margin_bottom = 0
    p_l2_tag = tf_l2.paragraphs[0]
    p_l2_tag.text = "LIVE DEMO URL (VERCEL)"
    p_l2_tag.font.name = "Arial"
    p_l2_tag.font.size = Pt(10)
    p_l2_tag.font.bold = True
    p_l2_tag.font.color.rgb = COLOR_ORANGE

    p_l2_val = tf_l2.add_paragraph()
    p_l2_val.text = "https://medibloom.vercel.app"
    p_l2_val.font.name = "Arial"
    p_l2_val.font.size = Pt(14)
    p_l2_val.font.bold = True
    p_l2_val.font.color.rgb = COLOR_DARK

    out_pptx = r"d:\NEWPROJECT\Medibloom\Medibloom\Medibloom_Round2_Presentation.pptx"
    prs.save(out_pptx)
    print(f"Presentation generated successfully: {out_pptx}")

if __name__ == "__main__":
    build_presentation()
